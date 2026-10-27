"""Multi-turn follow-ups that really need the conversation, the prompts, Word documents and the rag type."""
import pytest
from docx import Document
from fastapi.testclient import TestClient

from ragevalkit.test_cases.case_types import NON_RAG, RAG, is_non_rag
from ragevalkit.config import Config
from ragevalkit.documents.models import Chunk
from ragevalkit.documents.text_cleaning import normalise
from ragevalkit.documents.reader import read_document
from ragevalkit.documents.exceptions import DocumentError
from ragevalkit.generator import CASE_TYPE, Task, validate
from ragevalkit.generator.prompts import CATEGORY_RULES, build_messages, system_prompt
from ragevalkit.server import create_app

PASSAGE = ("Employees must submit expense claims within thirty days of travel. Managers approve claims above five hundred "
           "dollars within five working days. Late claims are paid in the next monthly cycle unless finance grants an "
           "exception in writing. Receipts are required for every item above twenty-five dollars.")


def chunk(text=PASSAGE):
    return Chunk(id="a#1", file="a.md", text=text, label="section: Expenses", pages=[(1, normalise(text))])


FIRST = {"question": "How long do I have to submit an expense claim?", "expected_answer": "Thirty days from the trip.",
         "evidence": ["Employees must submit expense claims within thirty days of travel."]}


def follow_up(question, depends="the thirty-day deadline from the first answer", quote="Late claims are paid in the next monthly cycle unless finance grants an exception in writing."):
    return {"question": question, "expected_answer": "They are paid in the next monthly cycle unless finance agrees.",
            "evidence": [quote], "depends_on": depends}


def check(*turns, avoid=()):
    return validate(Task("multiturn", "a.md", [chunk()], 1), {"turns": list(turns)}, turns=3, avoid=list(avoid))


# ---- a follow-up must need the conversation ----------------------------------------------------------

@pytest.mark.parametrize("question", [
    "And what happens if I miss that?",
    "What about late ones, are they still paid?",
    "Why is it so strict?",
    "In that case, who pays for the delay?",
    "What if it is a weekend?",
])
def test_follow_ups_that_refer_back_are_accepted_and_keep_what_they_rely_on(question):
    turns, reason = check(FIRST, follow_up(question))
    assert reason == "", reason
    assert turns[1]["needs_context"] == "the thirty-day deadline from the first answer" and "needs_context" not in turns[0]


@pytest.mark.parametrize("question", [
    "When are late expense claims paid by finance?",
    "What is required for each receipt item above twenty-five dollars?",
    "Who approves claims above five hundred dollars?",
])
def test_a_follow_up_that_could_be_asked_on_its_own_is_rejected_as_just_another_direct_question(question):
    turns, reason = check(FIRST, follow_up(question))
    assert turns is None and "could be asked on its own" in reason


def test_a_follow_up_that_restates_the_topic_instead_of_referring_back_is_rejected():
    restated = follow_up("And how long do I have to submit an expense claim late?")
    turns, reason = check(FIRST, restated)
    assert turns is None and "repeats the earlier question" in reason


@pytest.mark.parametrize("depends", ["", "x", None])
def test_every_follow_up_must_say_what_it_relies_on(depends):
    turns, reason = check(FIRST, follow_up("And what happens if I miss that?", depends=depends))
    assert turns is None and "depends_on" in reason


def test_the_first_turn_still_has_to_stand_alone_and_needs_no_depends_on():
    turns, _ = check(FIRST, follow_up("And what happens if I miss that?"))
    assert turns[0]["question"].startswith("How long")
    opener = dict(FIRST, question="And what about that?")           # a first question cannot lean on nothing... but it need not be checked for it
    assert check(opener, follow_up("And what happens if I miss that?"))[0] is not None
    assert check(dict(FIRST, question="According to the document, how long is the deadline?"), follow_up("And what if I miss that?"))[0] is None


def test_the_reason_given_back_to_the_model_tells_it_how_to_fix_the_follow_up():
    _, reason = check(FIRST, follow_up("When are late expense claims paid by finance?"))
    assert "refer" in reason and '"it"' in reason and "that" in reason


# ---- the prompts ------------------------------------------------------------------------------------

def test_the_multiturn_prompt_demands_follow_ups_that_cannot_stand_alone():
    prompt = system_prompt("multiturn", 3)
    for needle in ("exactly 3 turns", "IMPOSSIBLE to understand", "Do NOT repeat the topic", "depends_on", "Bad follow-up", "Good follow-up"):
        assert needle in prompt
    assert "{turns}" not in prompt


@pytest.mark.parametrize("category", ["direct", "indirect", "multipart", "multiturn"])
def test_every_prompt_states_the_grounding_rules_and_the_reply_format(category):
    prompt = system_prompt(category, 3)
    assert "Use ONLY the passages" in prompt and "Give numbers only" in prompt and "Reply with JSON only" in prompt
    assert category in CATEGORY_RULES


def test_the_prompt_shows_the_passages_with_their_labels_and_what_is_already_written():
    messages = build_messages(Task("multipart", "a.md", [chunk("Passage one text. " * 30), chunk("Passage two text. " * 30)], 1), 3, ["An earlier question?"])
    user = messages[1]["content"]
    assert "[Passage A] (section: Expenses)" in user and "[Passage B]" in user and "- An earlier question?" in user


# ---- Word documents ---------------------------------------------------------------------------------

def make_docx(path):
    doc = Document()
    doc.add_heading("Leave policy", level=1)
    for _ in range(5):
        doc.add_paragraph("Employees accrue one and a half days of paid leave for every month of service and may carry over five unused days into the next year with approval from their manager.")
    doc.add_heading("Absences", level=2)
    doc.add_paragraph("Sick leave longer than three consecutive days requires a medical certificate issued by a registered doctor.", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Leave type", "Notice needed"
    table.rows[1].cells[0].text, table.rows[1].cells[1].text = "Annual leave", "Fourteen days in advance"
    for _ in range(4):
        doc.add_paragraph("Peak periods require sign-off from the director, and emergency leave is handled by the people team on the same day as the request.")
    doc.save(str(path))


def test_a_word_document_is_read_with_its_headings_lists_and_tables(tmp_path):
    path = tmp_path / "leave.docx"
    make_docx(path)
    doc = read_document(path)
    assert doc.pages is None and doc.usable()
    text = "\n".join(c.text for c in doc.chunks)
    assert "carry over five unused days" in text and "- Sick leave longer than three" in text      # list item marked
    assert "Annual leave | Fourteen days in advance" in text                                       # table row kept together
    assert any(c.label == "section: Leave policy" for c in doc.chunks)


def test_unreadable_word_files_say_what_to_do(tmp_path):
    old = tmp_path / "old.doc"
    old.write_bytes(b"\xd0\xcf\x11\xe0 legacy word file")
    with pytest.raises(DocumentError, match="save it as .docx"):
        read_document(old)
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"this is not a zip file at all")
    with pytest.raises(DocumentError, match="could not be read as a Word document"):
        read_document(broken)
    empty = tmp_path / "empty.docx"
    Document().save(str(empty))
    with pytest.raises(DocumentError, match="almost no text"):
        read_document(empty)


@pytest.fixture
def client(tmp_path):
    cfg = Config(store_path=tmp_path / "s.json", require_login=False, eval_secret="s")
    with TestClient(create_app(cfg, fn=lambda q: "a", generation_llm=lambda m: "{}")) as c:
        yield c


def test_word_is_accepted_for_upload_and_listed_in_the_limits(client, tmp_path):
    assert "docx" in client.get("/api/generate/limits").json()["formats"]
    path = tmp_path / "leave.docx"
    make_docx(path)
    ok = client.post("/api/generate", files=[("files", ("leave.docx", path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"))],
                     data={"per_category": 1, "categories": '["direct"]'})
    assert ok.status_code == 202
    bad = client.post("/api/generate", files=[("files", ("notes.rtf", b"{\\rtf1}", "application/rtf"))])
    assert bad.status_code == 422 and "Word (.docx)" in bad.json()["detail"]


# ---- the type is rag or non-rag ----------------------------------------------------------------------

def test_the_type_is_rag_or_non_rag():
    assert (RAG, NON_RAG) == ("rag", "non-rag") and CASE_TYPE == RAG
    assert all(is_non_rag(v) for v in ("non-rag", "Non-RAG", "non_rag", "not rag", "NOT-RAG"))
    assert not any(is_non_rag(v) for v in ("rag", "query", "conversational", "", None))


def test_non_rag_cases_skip_retrieval_scoring_like_conversational_ones(client):
    manager = client.app.state.manager
    assert all(manager._is_conversational(t) for t in ("non-rag", "Non RAG", "not_rag", "conversational", "conv-chat"))
    assert not any(manager._is_conversational(t) for t in ("rag", "RAG", "query", "direct"))


def test_new_cases_default_to_rag_and_can_be_non_rag(client):
    made = client.post("/api/test-cases", json={"chat_id": "c1", "turn_no": 1, "question": "What is the refund window?", "expected_answer": "30 days"}).json()
    assert made["type"] == "rag"
    hello = client.post("/api/test-cases", json={"chat_id": "c2", "turn_no": 1, "question": "Hello there, how are you?", "expected_answer": "A friendly greeting", "type": "non-rag"}).json()
    assert hello["type"] == "non-rag"
    imported = client.post("/api/test-cases/import", files={"file": ("c.csv", b"question\nWhat is the SLA?\n", "text/csv")}).json()
    assert imported["new"] == 1
    assert {c["type"] for c in client.get("/api/test-cases").json()} == {"rag", "non-rag"}


# ---- evidence is cited by sentence number ------------------------------------------------------------

def numbered(text=PASSAGE):
    return chunk(text)


def test_the_prompt_shows_each_passage_as_numbered_sentences_to_cite():
    user = build_messages(Task("direct", "a.md", [numbered()], 1), 3, [])[1]["content"]
    assert "[A1] Employees must submit expense claims within thirty days of travel." in user
    assert "[A4] Receipts are required for every item above twenty-five dollars." in user


@pytest.mark.parametrize("evidence", [["A2"], ["[A2]"], "A2", ["a2"], [2], ["2"], [{"id": "A2"}]])
def test_evidence_cited_by_number_is_stored_as_the_passages_own_sentence(evidence):
    turns, reason = validate(Task("direct", "a.md", [numbered()], 1), dict(FIRST, evidence=evidence), turns=3, avoid=[])
    assert reason == "", reason
    assert [e["quote"] for e in turns[0]["evidence"]] == ["Managers approve claims above five hundred dollars within five working days."]
    assert turns[0]["evidence"][0]["passage"] == "A" and turns[0]["evidence"][0]["page"] == "page 1"


def test_several_numbers_and_ranges_are_read_and_unknown_numbers_prove_nothing():
    task = Task("direct", "a.md", [numbered()], 1)
    turns, _ = validate(task, dict(FIRST, evidence=["A1, A3"]), turns=3, avoid=[])
    assert len(turns[0]["evidence"]) == 2
    turns, _ = validate(task, dict(FIRST, evidence=["A1-A3"]), turns=3, avoid=[])
    assert len(turns[0]["evidence"]) == 3
    turns, reason = validate(task, dict(FIRST, evidence=["A40"]), turns=3, avoid=[])
    assert turns is None and "sentence numbers" in reason


def test_multipart_numbers_point_at_the_right_passage():
    two = Task("multipart", "a.md", [numbered(), numbered("Visitors must sign in at the front desk before entering the building. Badges must be worn at all times.")], 1)
    turns, reason = validate(two, dict(FIRST, question="What is the claim deadline, and what must visitors do?", evidence=["A1", "B1"]), turns=3, avoid=[])
    assert reason == "" and [e["passage"] for e in turns[0]["evidence"]] == ["A", "B"]
    assert "front desk" in turns[0]["evidence"][1]["quote"]


def test_a_wrapped_pdf_sentence_is_one_sentence_and_list_items_stay_apart():
    text = ("Employees must submit expense claims\n\nwithin thirty days of travel. Managers approve\n\nclaims above five hundred dollars.\n\n"
            "- Receipts are needed above twenty-five dollars\n- Late claims wait for the next cycle")
    sentences = numbered(text).sentences
    assert sentences[0] == "Employees must submit expense claims within thirty days of travel."
    assert sentences[1] == "Managers approve claims above five hundred dollars."
    assert sentences[2:] == ["Receipts are needed above twenty-five dollars", "Late claims wait for the next cycle"]


def test_a_quote_is_still_accepted_when_the_model_copies_instead_of_citing():
    turns, reason = validate(Task("direct", "a.md", [numbered()], 1), FIRST, turns=3, avoid=[])
    assert reason == "" and turns[0]["evidence"][0]["quote"].startswith("Employees must submit")


# ---- conversations: more ways to refer back, and the reason names the turn ---------------------------

@pytest.mark.parametrize("question", [
    "Does this apply on weekends too?",
    "Is there another way to get paid sooner?",
    "What happens to the rest of them?",
    "Do both need approval?",
])
def test_more_everyday_ways_of_referring_back_are_accepted(question):
    assert check(FIRST, follow_up(question))[1] == ""


def test_a_rejected_follow_up_names_its_turn_and_quotes_the_question_back():
    _, reason = check(FIRST, follow_up("When are late expense claims paid by finance?"))
    assert reason.startswith("turn 2:") and '"When are late expense claims paid by finance?"' in reason
    _, reason = check(FIRST, follow_up("And what if I miss that?", quote="Nothing like this is in the passage at all."))
    assert reason.startswith("turn 2:") and "sentence numbers" in reason


def test_single_question_reasons_have_no_turn_prefix():
    _, reason = validate(Task("direct", "a.md", [chunk()], 1), dict(FIRST, question="Why?"), turns=3, avoid=[])
    assert not reason.startswith("turn")


def test_markdown_headings_lists_tables_and_abbreviations_are_numbered_sensibly():
    text = ("## 3. Retell Caller Agent\n\nThe agent answers inbound calls and routes them. It supports **English** and Spanish.\n\n"
            "### Escalation rules\n- Calls longer than 10 minutes are escalated to a human agent\n\n"
            "| Plan | Minutes |\n|---|---|\n| Basic | 500 |\n\n"
            "Billing is monthly. Overage is charged at $0.10 per minute, e.g. Basic plans pay more.")
    assert chunk(text).sentences == [
        "Retell Caller Agent", "The agent answers inbound calls and routes them.", "It supports English and Spanish.",
        "Escalation rules", "Calls longer than 10 minutes are escalated to a human agent",
        "Plan Minutes Basic 500", "Billing is monthly.", "Overage is charged at $0.10 per minute, e.g. Basic plans pay more.",
    ]
