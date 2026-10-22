"""Reading test cases from a spreadsheet, paste or CSV — forgivingly, but never silently.

The importer works in two stages so the dashboard can show a preview before anything is saved:

1. ``read_table`` + ``detect_mapping`` find the data and work out which column is which (by common
   names, the developer's ``column_map``, or a choice made in the dialog).
2. ``parse_rows`` turns the table into clean rows and, for every row it cannot use, says which row
   and why. Nothing is quietly "fixed": a turn of "abc" is a problem, not a 1.

What happens to the clean rows (new / updated / unchanged) is decided by ``EvalStore.import_cases``,
which uses the same code for the preview (``dry_run=True``) and the real import, so they never disagree.
"""
from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import asdict, dataclass
from pathlib import Path

import pandas as pd

from ragevalkit.exceptions import DatasetError

MAX_ROWS = 20_000
MAX_BYTES = 10 * 1024 * 1024

FIELDS = ("chat_id", "turn_no", "type", "question", "expected_answer", "expected_sources", "groups")

FIELD_LABELS = {
    "chat_id": "Chat ID",
    "turn_no": "Turn number",
    "type": "Type",
    "question": "Question",
    "expected_answer": "Expected answer",
    "expected_sources": "Expected sources",
    "groups": "Groups",
}

# Header names we recognise for each field, compared after lower-casing and turning punctuation into "_".
ALIASES: dict[str, tuple[str, ...]] = {
    "chat_id": ("chat_id", "chat", "chatid", "conversation_id", "conversation", "session_id", "session", "thread_id", "thread"),
    "turn_no": ("turn_no", "turn", "turn_number", "turn_num", "turnno", "step", "message_no"),
    "type": ("type", "question_type", "category", "kind"),
    "question": ("question", "query", "user_query", "user_question", "prompt", "input", "user_input", "utterance"),
    "expected_answer": ("expected_answer", "expected", "answer", "gold_answer", "reference_answer", "reference", "ground_truth", "expected_response"),
    "expected_sources": ("expected_sources", "expected_source", "sources", "source", "documents", "expected_documents", "references"),
    "groups": ("groups", "group", "tags", "tag", "labels", "label"),
}

_TEXT_SUFFIXES = {".csv", ".tsv", ".txt"}
_EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
SUPPORTED_SUFFIXES = _TEXT_SUFFIXES | _EXCEL_SUFFIXES


def _norm(header) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(header).strip().lower()).strip("_")


def _s(value) -> str:
    """A cell as clean text: empty for blanks and NaN, no stray surrounding space."""
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    text = str(value).strip()
    return "" if text.lower() == "nan" else text


def read_table(path: Path, sheet: str | None = None) -> tuple[pd.DataFrame, list[str], str | None]:
    """Read a spreadsheet or delimited text file with every cell as text.

    Returns ``(table, sheet_names, sheet_used)``; the last two are empty / None for text files.
    Raises DatasetError for a file that cannot be read.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DatasetError(f"Unsupported file type '{suffix or 'unknown'}'. Use .xlsx, .csv or .tsv.")
    try:
        if suffix in _TEXT_SUFFIXES:
            sample = path.read_text(encoding="utf-8-sig", errors="replace")[:8192]
            try:
                sep = csv.Sniffer().sniff(sample, delimiters=",;	|").delimiter
            except csv.Error:
                sep = ","  # e.g. a single column of questions
            df = pd.read_csv(path, dtype=str, keep_default_na=False, sep=sep, encoding="utf-8-sig")
            sheets, used = [], None
        else:
            with pd.ExcelFile(path, engine="openpyxl") as xl:
                sheets = list(xl.sheet_names)
                used = sheet if sheet in sheets else ("dataset" if "dataset" in sheets else sheets[0])
                df = xl.parse(used, dtype=str, keep_default_na=False)
    except DatasetError:
        raise
    except Exception as e:
        raise DatasetError(f"Could not read file: {e}") from e

    df.columns = [str(c).strip() for c in df.columns]
    # drop spreadsheet filler: columns with no header and no data
    keep = [c for c in df.columns if not (c.lower().startswith("unnamed:") and not df[c].map(_s).any())]
    df = df[keep]
    if len(df) > MAX_ROWS:
        raise DatasetError(f"The file has {len(df):,} rows; the limit is {MAX_ROWS:,}. Split it into smaller files.")
    return df, sheets, used


def detect_mapping(columns: list[str], column_map: dict[str, str] | None = None) -> dict[str, str | None]:
    """Which file column holds each field: ``{field: column or None}``.

    Order of trust: the developer's ``serve(column_map=...)``, then well-known header names.
    """
    by_norm: dict[str, str] = {}
    for col in columns:
        by_norm.setdefault(_norm(col), col)
    mapping: dict[str, str | None] = {f: None for f in FIELDS}

    for source, target in (column_map or {}).items():
        field, col = str(target).strip().lower(), by_norm.get(_norm(source))
        if field in mapping and col and mapping[field] is None and col not in mapping.values():
            mapping[field] = col

    for field in FIELDS:
        if mapping[field] is not None:
            continue
        for alias in ALIASES[field]:
            col = by_norm.get(alias)
            if col and col not in mapping.values():
                mapping[field] = col
                break
    return mapping


def clean_mapping(columns: list[str], requested: dict | None, detected: dict[str, str | None]) -> dict[str, str | None]:
    """Apply the user's choices from the dialog on top of the detected mapping (unknown columns ignored)."""
    if not requested:
        return detected
    mapping = dict(detected)
    for field in FIELDS:
        if field in requested:
            col = requested[field]
            mapping[field] = col if col in columns else None
    return mapping


@dataclass
class Problem:
    row: int          # the row number as the user sees it in the spreadsheet (header is row 1)
    reason: str
    cells: dict[str, str]

    def as_dict(self) -> dict:
        return asdict(self)


def split_labels(text: str) -> list[str]:
    """'billing; refunds | faq' -> ['billing', 'refunds', 'faq'] (case-insensitive de-dupe, order kept)."""
    seen: dict[str, str] = {}
    for part in re.split(r"[;|,\n]", text or ""):
        name = part.strip()
        if name:
            seen.setdefault(name.lower(), name[:80])
    return list(seen.values())


def parse_rows(df: pd.DataFrame, mapping: dict[str, str | None]) -> tuple[list[dict], list[Problem]]:
    """Turn a table into clean test-case rows plus a list of the rows that could not be used.

    Missing columns are filled in sensibly rather than rejected: no chat column makes every question its
    own single-turn chat, no turn column numbers the turns of each chat in file order, a blank type
    becomes "query". Anything that is *present but wrong* is reported, never altered.
    """
    cols = {f: mapping.get(f) for f in FIELDS}
    rows: list[dict] = []
    problems: list[Problem] = []
    seen_key: dict[tuple, int] = {}
    seen_slot: dict[tuple, int] = {}
    next_turn: dict[str, int] = {}

    for index, record in enumerate(df.to_dict("records")):
        number = index + 2
        cells = {str(k): _s(v) for k, v in record.items()}
        if not any(cells.values()):
            continue  # a blank line

        def get(field: str) -> str:
            col = cols[field]
            return cells.get(col, "") if col else ""

        question = get("question")
        if not question:
            problems.append(Problem(number, "Question is empty", cells))
            continue

        if cols["chat_id"] is None:
            chat = "q-" + hashlib.sha1(question.encode("utf-8")).hexdigest()[:6]
        else:
            chat = get("chat_id")
            if not chat:
                problems.append(Problem(number, "Chat ID is empty", cells))
                continue

        if cols["turn_no"] is None:
            if cols["chat_id"] is None:
                turn = 1  # every question is its own chat, so a repeated question is a duplicate, not turn 2
            else:
                next_turn[chat] = next_turn.get(chat, 0) + 1
                turn = next_turn[chat]
        else:
            raw = get("turn_no")
            if raw == "":
                turn = 1
            else:
                try:
                    value = float(raw)
                    if value != int(value) or value < 1:
                        raise ValueError
                    turn = int(value)
                except ValueError:
                    problems.append(Problem(number, f"Turn number '{raw}' is not a whole number, 1 or higher", cells))
                    continue

        key, slot = (chat, turn, question), (chat, turn)
        if key in seen_key:
            where = "" if cols["chat_id"] is None else f" (chat {chat}, turn {turn})"
            problems.append(Problem(number, f"Same question as row {seen_key[key]}{where}", cells))
            continue
        if slot in seen_slot:
            problems.append(Problem(
                number, f"Chat {chat} turn {turn} is already used by row {seen_slot[slot]} with a different question", cells))
            continue
        seen_key[key], seen_slot[slot] = number, number

        rows.append({
            "row": number,
            "chat_id": chat,
            "turn_no": turn,
            "type": get("type") or "query",
            "question": question,
            "expected_answer": get("expected_answer"),
            "expected_sources": get("expected_sources"),
            "groups": split_labels(get("groups")),
        })
    return rows, problems


def template_workbook() -> bytes:
    """A ready-to-fill .xlsx: the right headers, two example rows and a short guide."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "dataset"
    headers = ["chat_id", "turn_no", "type", "question", "expected_answer", "expected_sources", "groups"]
    ws.append(headers)
    ws.append(["refund-1", 1, "query", "What is the refund window?", "Refunds are possible within 30 days.", "refund_policy.pdf", "Billing; Refunds"])
    ws.append(["refund-1", 2, "query", "And for annual plans?", "Annual plans have a 60-day window.", "refund_policy.pdf , annual_terms.pdf", "Billing"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFG", (14, 10, 12, 44, 44, 34, 22)):
        ws.column_dimensions[letter].width = width

    guide = wb.create_sheet("How to fill")
    for line in [
        ("Column", "What to put there"),
        ("question", "Required. The question to ask the assistant."),
        ("expected_answer", "The answer you expect. Used to score the assistant's reply."),
        ("chat_id", "Questions that belong to one conversation share a chat_id. Leave the column out and every question becomes its own chat."),
        ("turn_no", "1 for the first message of a chat, 2 for the follow-up, and so on. Leave it out to number them in file order."),
        ("type", "Optional label such as query or conversational. Blank becomes 'query'."),
        ("expected_sources", "Optional. Documents the assistant should retrieve. Use | between documents that are all required and , between acceptable alternatives."),
        ("groups", "Optional. One or more group names separated by ; — they are created if they do not exist yet."),
        ("", ""),
        ("Re-importing", "Importing the same file again changes nothing. Edited answers update the existing test case as a new version."),
    ]:
        guide.append(list(line))
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
