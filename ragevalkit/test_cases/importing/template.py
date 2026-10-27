"""The downloadable template with a short guide."""




def template_workbook() -> bytes:
    """A ready-to-fill .xlsx: the right headers, two example rows and a short guide."""
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "dataset"
    headers = ["chat_id", "turn_no", "type", "question", "expected_answer", "expected_sources", "groups"]
    ws.append(headers)
    ws.append(["refund-1", 1, "rag", "What is the refund window?", "Refunds are possible within 30 days.", "refund_policy.pdf", "Billing; Refunds"])
    ws.append(["refund-1", 2, "rag", "And for annual plans?", "Annual plans have a 60-day window.", "refund_policy.pdf , annual_terms.pdf", "Billing"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCDEFG", (14, 10, 12, 44, 44, 34, 22)):
        ws.column_dimensions[letter].width = width

    guide = wb.create_sheet("How to fill")
    for line in [
        ("Column", "What to put there"),
        ("question", "Required. The question to ask the assistant."),
        ("expected_answer", "The answer you expect. Used to score the assistant's reply."),
        ("chat_id", "Questions that belong to one conversation share a chat_id. Leave the column out and every question becomes its own chat."),
        ("turn_no", "1 for the first message of a chat, 2 for the follow-up, and so on. Leave it out to number them in file order."),
        ("type", "rag (the answer comes from your documents) or non-rag (a greeting or small talk, where no retrieval is expected). Blank means rag."),
        ("expected_sources", "Optional. Documents the assistant should retrieve. Use | between documents that are all required and , between acceptable alternatives."),
        ("groups", "Optional. One or more group names separated by ; — they are created if they do not exist yet."),
        ("", ""),
        ("Re-importing", "Importing the same file again changes nothing. Edited answers update the existing test case as a new version."),
    ]:
        guide.append(list(line))
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 110
    for cell in guide[1]:
        cell.font = Font(bold=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
